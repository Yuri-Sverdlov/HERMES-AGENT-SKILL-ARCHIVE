# -*- coding: utf-8 -*-
"""
Overview-projects-on-computer — обновление адресов проектов в базе знаний.

Назначение: когда пользователь перемещает проекты между папками _MY_PROGRAMMING*,
этот скрипт пересканирует верхнюю папку и обновляет ТОЛЬКО пути (dirs) в базе.
Метаданные (статус, цель, функционал, GitHub-репо) сохраняются.

Использование:
    python update_projects.py "F:/"

Что делает:
    1. Читает references/projects.json (ключ = нормализованное имя проекта).
    2. Сканирует верхнюю папку, находит подпапки _MY_PROGRAMMING* и проекты внутри.
    3. Сопоставляет по имени проекта (регистронезависимо, _ и - эквивалентны).
    4. Обновляет поле dirs; новые проекты добавляет; исчезнувшие помечает "НЕ НАЙДЕН".
    5. Перезаписывает projects.json и перегенерирует README + сирот + Word-файлы.
"""
import os, sys, re, json, datetime

# ---------- конфигурация ----------
OUTPUT_DIR = r"C:/Users/Yuri/Documents/PROJECTS_KNOWLEDGE_BASE"
SKILL_DIR = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = os.path.join(SKILL_DIR, "..", "references", "projects.json")

STATUS_ICON = {"живой": "🟢", "спящий": "🟡", "мёртвый": "⚫", "окончен": "🔵", "новый": "🆕"}

def norm(name):
    """Нормализация имени: регистр не важен, _ - пробелы эквивалентны."""
    return re.sub(r"[\s_\-]+", "-", name).upper().strip("-")

def load_json():
    with open(JSON_PATH, encoding="utf-8") as f:
        return json.load(f)

def save_json(data):
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

# ---------- сканирование ----------
def scan_myprogramming(top):
    """Найти все подпапки _MY_PROGRAMM* (включая опечатку _MY_PROGRAMSING)
    и проекты внутри них. Возвращает {norm_name: [полные пути]}."""
    result = {}
    if not os.path.isdir(top):
        return result
    for entry in sorted(os.listdir(top)):
        if not entry.startswith("_MY_PROGRAM"):
            continue
        dpath = os.path.join(top, entry)
        if not os.path.isdir(dpath):
            continue
        for sub in sorted(os.listdir(dpath)):
            if sub.startswith(".") or sub.startswith("~$"):
                continue
            fpath = os.path.join(dpath, sub)
            if os.path.isdir(fpath):
                # нормализовать к прямым слэшам (переносимость между машинами)
                fpath = fpath.replace("\\", "/")
                result.setdefault(norm(sub), []).append(fpath)
    return result

def detect_new(fpath):
    """Лёгкая авто-детекция нового проекта."""
    has_git = os.path.isdir(os.path.join(fpath, ".git"))
    langs = set()
    for root, dirs, files in os.walk(fpath):
        dirs[:] = [d for d in dirs if d not in (".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build")]
        for fn in files:
            e = os.path.splitext(fn)[1].lower()
            if e in (".py", ".js", ".ts", ".html", ".css", ".json", ".md", ".ipynb", ".mjs", ".tsx", ".jsx"):
                langs.add(e)
        if len(files) > 500:
            break
    mtime = datetime.datetime.fromtimestamp(os.path.getmtime(fpath))
    age_months = (datetime.datetime.now() - mtime).days / 30.0
    status = "мёртвый" if age_months > 6 else ("спящий" if age_months > 3 else "живой")
    # попробовать вытащить первую строку README как purpose
    purpose = ""
    for cand in ("README.md", "readme.md", "README.MD"):
        rp = os.path.join(fpath, cand)
        if os.path.isfile(rp):
            try:
                for line in open(rp, encoding="utf-8", errors="ignore"):
                    s = line.strip().lstrip("# ").strip()
                    if s and len(s) > 3:
                        purpose = s[:80]
                        break
            except Exception:
                pass
            break
    return status, purpose or "—", ",".join(sorted(langs)) or None

def is_myprogramming(path, top):
    """Путь лежит под _MY_PROGRAMMING* внутри top?"""
    p = path.replace("\\", "/").rstrip("/")
    t = top.replace("\\", "/").rstrip("/")
    if p == t:
        return False
    rel = p[len(t):].lstrip("/") if p.startswith(t + "/") else p
    return rel.startswith("_MY_PROGRAM")

# ---------- основная логика ----------
def update(top):
    data = load_json()
    projects = data["projects"]
    scan = scan_myprogramming(top)

    changed = []
    added = []
    missing = []

    # 1. Обновить существующие
    for key, p in projects.items():
        # отделить "корневые" пути (вне _MY_PROGRAMMING) — их не трогаем
        non_mp = [d for d in p["dirs"] if not is_myprogramming(d, top)]
        had_mp = any(is_myprogramming(d, top) for d in p["dirs"])
        fresh = scan.get(key, [])
        new_dirs = non_mp + fresh
        if new_dirs != p["dirs"]:
            changed.append((p["name"], p["dirs"], new_dirs))
        p["dirs"] = new_dirs
        if had_mp and not fresh:
            missing.append(p["name"])
            if "НЕ НАЙДЕН" not in p["note"]:
                p["note"] = (p["note"] + " | НЕ НАЙДЕН в " + top).strip(" |")

    # 2. Добавить новые проекты
    for key, paths in scan.items():
        if key in projects:
            continue
        fpath = paths[0]
        status, purpose, lang = detect_new(fpath)
        display_name = os.path.basename(fpath)
        projects[key] = {
            "name": display_name,
            "repo": None,
            "status": status,
            "purpose": purpose,
            "func": "",
            "lang": lang,
            "note": "авто-добавлен " + datetime.date.today().isoformat(),
            "dirs": paths,
        }
        added.append(display_name)

    data["meta"]["top_folder"] = top
    data["meta"]["updated"] = datetime.date.today().isoformat()
    save_json(data)

    return data, changed, added, missing

# ---------- генерация выходных файлов ----------
def render(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    projects = data["projects"]

    def sortkey(p):
        return p["name"].lower()

    # сгруппировать по "отображаемой директории"
    top = data["meta"]["top_folder"]
    groups = {}  # display_dir -> [projects]
    for p in projects.values():
        if p["dirs"]:
            d = p["dirs"][0]
            if is_myprogramming(d, top):
                disp = os.path.basename(os.path.dirname(d))
            else:
                disp = top + " (корень)"
        else:
            disp = "только GitHub"
        groups.setdefault(disp, []).append(p)

    order = sorted(groups.keys(), key=lambda g: (g == "только GitHub", g == top + " (корень)", g.lower()))

    # --- README.md ---
    md = ["# База знаний проектов (обновлено " + data["meta"]["updated"] + ")\n"]
    md.append("> Верхняя папка: `" + top + "`. Статусы: 🟢 живой, 🟡 спящий, ⚫ мёртвый, 🆕 новый.\n")
    total = len(projects)
    n_orph = sum(1 for p in projects.values() if not p["repo"])
    md.append(f"- Всего проектов: **{total}** (в т.ч. сирот без GitHub: **{n_orph}**)\n")
    for disp in order:
        projs = groups[disp]
        md.append(f"\n## 📁 {disp}\n")
        md.append("| Проект | Статус | GitHub | Язык | Цель |")
        md.append("|---|---|---|---|---|")
        for p in sorted(projs, key=sortkey):
            icon = STATUS_ICON.get(p["status"], "❔")
            repo = f'[{p["repo"]}](https://github.com/{data["meta"]["github_user"]}/{p["repo"]})' if p["repo"] else "—"
            md.append(f"| {icon} {p['name']} | {p['status']} | {repo} | {p.get('lang') or '—'} | {(p.get('purpose') or '')[:60]} |")
    with open(os.path.join(OUTPUT_DIR, "README_проекты.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    # --- Сироты ---
    omd = ["# Проекты без GitHub («сироты»)\n"]
    omd.append(f"> Обновлено {data['meta']['updated']}. Всего: **{n_orph}**.\n")
    omd.append("| Директория | Проект | Статус | Функционал |")
    omd.append("|---|---|---|---|")
    for disp in order:
        for p in sorted(groups[disp], key=sortkey):
            if p["repo"]:
                continue
            icon = STATUS_ICON.get(p["status"], "❔")
            omd.append(f"| {disp} | {icon} {p['name']} | {p['status']} | {(p.get('func') or '—')[:70]} |")
    with open(os.path.join(OUTPUT_DIR, "СИРОТЫ_не_на_GitHub.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(omd))

    # --- Word на каждую директорию ---
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return

    for disp in order:
        if disp == "только GitHub":
            continue
        projs = sorted(groups[disp], key=sortkey)
        doc = Document()
        sec = doc.sections[0]
        sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(1.5)
        st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
        rpr = st.element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
        rf.set(qn("w:cs"), "Times New Roman"); rf.set(qn("w:eastAsia"), "Times New Roman")

        t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr = t.add_run(f"Проекты: {disp}"); tr.bold = True; tr.font.size = Pt(17)
        sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sub.add_run(f"Подробный реестр · {len(projs)} проектов · {data['meta']['updated']}")
        sr.italic = True; sr.font.size = Pt(10); sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        for p in projs:
            doc.add_paragraph()
            h = doc.add_paragraph()
            icon = STATUS_ICON.get(p["status"], "❔")
            hr = h.add_run(f"{icon} {p['name']}")
            hr.bold = True; hr.font.size = Pt(13); hr.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)

            def line(label, value):
                if not value:
                    return
                pp = doc.add_paragraph()
                r1 = pp.add_run(label + ": "); r1.bold = True; r1.font.size = Pt(11)
                r2 = pp.add_run(value); r2.font.size = Pt(11)
                pp.paragraph_format.left_indent = Cm(0.5)

            line("Статус", p["status"])
            line("Путь", " ; ".join(p["dirs"]) if p["dirs"] else "—")
            line("GitHub", f"https://github.com/{data['meta']['github_user']}/{p['repo']}" if p["repo"] else "нет (сирота)")
            line("Язык/стек", p.get("lang") or "—")
            line("Цель создания", p.get("purpose"))
            line("Переиспользуемый функционал", p.get("func"))
            line("Примечание", p.get("note"))

        safe = disp.replace("/", "_").replace(":", "").replace(" ", "_").strip("_")
        doc.save(os.path.join(OUTPUT_DIR, f"Word_{safe}.docx"))

# ---------- CLI ----------
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python update_projects.py <ВЕРХНЯЯ_ПАПКА>")
        sys.exit(1)
    top = sys.argv[1]
    top = top.replace("\\", "/").rstrip("/")
    data, changed, added, missing = update(top)
    render(data)
    print("=" * 60)
    print(f"Верхняя папка: {top}")
    print(f"Проектов в базе: {len(data['projects'])}")
    print(f"Путей обновлено: {len(changed)}")
    for name, old, new in changed:
        print(f"   ✏️ {name}: {old} → {new}")
    print(f"Новых проектов: {len(added)}")
    for name in added:
        print(f"   🆕 {name}")
    print(f"Не найдено (были в _MY_PROGRAMMING): {len(missing)}")
    for name in missing:
        print(f"   ⚠️ {name}")
    print(f"Файлы перегенерированы в: {OUTPUT_DIR}")
