#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Шаблон: создание Word-документа из распознанного текста.
Замените BLOCK_* на реальный текст, OUTPUT_PATH на путь сохранения."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import os

OUTPUT_PATH = r"G:\path\to\output.docx"
TITLE = "НАЗВАНИЕ ДОКУМЕНТА"
SUBTITLE = "Подзаголовок"
SOURCE = "ИСТОЧНИК"

doc = Document()

# --- Page setup: A4 landscape ---
section = doc.sections[0]
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.orientation = WD_ORIENT.LANDSCAPE
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 51, 102)
    return h


def add_para(text, bold=False, italic=False, align=None, size=12, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return p


def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p


# ========================
# TITLE PAGE
# ========================
doc.add_paragraph()
add_para(SOURCE, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
         color=RGBColor(0, 51, 102), space_after=4)
add_para(SUBTITLE, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
         color=RGBColor(100, 100, 100), space_after=16)
add_para(TITLE, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
add_para("Распознанный текст скриншота • Август 2026", italic=True, size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(128, 128, 128))
doc.add_page_break()

# ========================
# БЛОК 1
# ========================
add_heading("Блок 1. Название блока", level=1)
add_body("BLOCK_1_TEXT_HERE")

# ========================
# БЛОК 2
# ========================
add_heading("Блок 2. Название блока", level=1)
add_body("BLOCK_2_TEXT_HERE")

# ... повторить для блоков 3–5 ...

# ========================
# МЕТРИКИ
# ========================
doc.add_page_break()
add_heading("Метрики публикаций", level=2)
add_para("Пост 1:", bold=True, size=12, space_after=4)
add_para("❤️ N  •  🔥 N  •  👎 N  •  🤔 N  •  👍 N", size=11, space_after=2)
add_para("👁 N просмотров  •  HH:MM", size=11, color=RGBColor(128, 128, 128))

# ========================
# SAVE
# ========================
doc.save(OUTPUT_PATH)
print(f"DOCX saved: {OUTPUT_PATH}")
print(f"File size: {os.path.getsize(OUTPUT_PATH)} bytes")